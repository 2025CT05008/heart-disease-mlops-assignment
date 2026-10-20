"""Create the student-facing macOS and VS Code execution guide."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Mac_VS_Code_Execution_Guide.docx"
PROJECT = "/Users/bikash.bishoyi/Documents/Codex/2026-10-18/as-an-ai-ml-student-i/outputs/heart-disease-mlops"


def shade(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def code(document: Document, command: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    shade(paragraph, "F2F2F2")
    run = paragraph.add_run(command)
    run.font.name = "Menlo"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Menlo")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Menlo")
    run.font.size = Pt(9)


def section(document: Document, title: str) -> None:
    document.add_heading(title, level=1)


def bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def main() -> None:
    d = Document()
    s = d.sections[0]
    s.top_margin = Inches(0.7)
    s.bottom_margin = Inches(0.7)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)
    normal = d.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        style = d.styles[name]
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = d.add_heading("Mac VS Code Execution Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    sub = d.add_paragraph("Heart Disease MLOps Assignment")
    sub.style = d.styles["Subtitle"]
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(
        "This guide shows how to run the complete assignment locally on macOS using Visual Studio Code. Follow the steps in order. Keep the API, MLflow, and Docker terminals separate when those services are running."
    )

    section(d, "1 Prerequisites")
    d.add_paragraph("Install these before starting:")
    bullets(d, [
        "Visual Studio Code with the Microsoft Python extension.",
        "Python 3.11 or Python 3.12. Check with python3 --version.",
        "Optional for container work: Docker Desktop for Mac. Start Docker Desktop before the Docker section.",
        "Optional for Kubernetes: enable Kubernetes in Docker Desktop or install Minikube and kubectl.",
    ])
    d.add_paragraph("The project folder used in this guide is:")
    code(d, PROJECT)

    section(d, "2 Open the Project in VS Code")
    d.add_paragraph("Open Terminal on your Mac and run the following commands. The second command opens the current folder in VS Code. If the code command is unavailable, use File then Open Folder in VS Code and select the same location.")
    code(d, f'cd "{PROJECT}"\ncode .')
    d.add_paragraph("In VS Code, open Terminal then New Terminal. All commands below are entered in the integrated terminal unless a separate terminal is specified.")

    section(d, "3 Create and Select the Python Environment")
    d.add_paragraph("Create a project-local virtual environment and install the pinned dependencies. Run this once for a new clone or copied project folder.")
    code(d, "python3 -m venv .venv\nsource .venv/bin/activate\npython -m pip install --upgrade pip\npip install -r requirements.txt")
    d.add_paragraph("Then select the environment in VS Code:")
    bullets(d, [
        "Press Cmd Shift P.",
        "Choose Python Select Interpreter.",
        "Select the interpreter ending in .venv/bin/python.",
    ])
    d.add_paragraph("Whenever you open a new terminal, activate the environment again:")
    code(d, "source .venv/bin/activate\nexport PYTHONPATH=src")

    section(d, "4 Prepare Data Train the Model and Validate")
    d.add_paragraph("Run these commands from the project root. Data preparation creates the cleaned CSV and EDA figures. Training compares Logistic Regression and Random Forest, saves the selected pipeline, and records MLflow runs.")
    code(d, "export PYTHONPATH=src\npython scripts/prepare_data.py\npython -m heart_ml.train --data data/cleaned/heart_disease_cleaned.csv")
    d.add_paragraph("Validate code quality and core behavior after training:")
    code(d, "pytest -q\nruff check src tests scripts")
    d.add_paragraph("Expected files after success:")
    bullets(d, [
        "data/cleaned/heart_disease_cleaned.csv",
        "reports/figures/ with EDA charts",
        "reports/metrics.json with the two-model comparison",
        "models/heart_disease_pipeline.joblib",
        "mlruns/ containing MLflow experiment data",
    ])

    section(d, "5 View MLflow Experiments")
    d.add_paragraph("Open a second VS Code terminal, activate the virtual environment, and start MLflow. Leave this terminal running while you inspect results.")
    code(d, "cd \"" + PROJECT + "\"\nsource .venv/bin/activate\nmlflow ui --backend-store-uri ./mlruns")
    d.add_paragraph("Open http://127.0.0.1:5000 in a browser. Compare both runs using ROC-AUC, accuracy, precision, recall, model name, and random seed.")

    section(d, "6 Run and Test the FastAPI Service")
    d.add_paragraph("Open a third terminal. Start the API only after model training has created models/heart_disease_pipeline.joblib.")
    code(d, "cd \"" + PROJECT + "\"\nsource .venv/bin/activate\nexport PYTHONPATH=src\nuvicorn heart_ml.api:app --host 0.0.0.0 --port 8000")
    d.add_paragraph("Open http://127.0.0.1:8000/docs. In the POST /predict endpoint, click Try it out and use the sample JSON below.")
    code(d, '{\n  "age": 63, "sex": 1, "cp": 1, "trestbps": 145, "chol": 233,\n  "fbs": 1, "restecg": 2, "thalach": 150, "exang": 0,\n  "oldpeak": 2.3, "slope": 3, "ca": 0, "thal": 6\n}')
    d.add_paragraph("You can also test directly from a fourth terminal:")
    code(d, "curl http://127.0.0.1:8000/health\ncurl http://127.0.0.1:8000/metrics")

    section(d, "7 Run the Docker Container and Monitoring")
    d.add_paragraph("Confirm Docker Desktop shows Engine running. Build and start the API container:")
    code(d, "docker build -t heart-disease-api:local .\ndocker run --rm -p 8000:8000 heart-disease-api:local")
    d.add_paragraph("Alternatively, start the API and Prometheus together:")
    code(d, "docker compose up --build")
    bullets(d, [
        "API documentation: http://127.0.0.1:8000/docs",
        "Prometheus: http://127.0.0.1:9090",
        "In Prometheus, query heart_predictions_total after making one or more API requests.",
    ])

    section(d, "8 Optional Local Kubernetes Test")
    d.add_paragraph("First update the image placeholder in k8s/deployment.yaml to an image available to your cluster. For Docker Desktop Kubernetes, a local image can be used when imagePullPolicy is IfNotPresent. Apply the manifests:")
    code(d, "kubectl apply -f k8s/namespace.yaml\nkubectl apply -f k8s/deployment.yaml\nkubectl apply -f k8s/service.yaml\nkubectl apply -f k8s/hpa.yaml\nkubectl -n heart-ml get pods,svc,hpa\nkubectl -n heart-ml port-forward service/heart-disease-api 8000:80")
    d.add_paragraph("Use the Swagger UI at http://127.0.0.1:8000/docs while port-forwarding is active.")

    section(d, "9 Troubleshooting")
    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Problem"
    table.rows[0].cells[1].text = "Resolution"
    for problem, resolution in [
        ("Module not found", "Activate .venv, run pip install -r requirements.txt, then export PYTHONPATH=src."),
        ("Model file not found", "Run the data preparation and training commands before starting Uvicorn."),
        ("Port already in use", "Stop the old process with Ctrl C or use a different port such as 8001."),
        ("Docker command unavailable", "Install and start Docker Desktop, then reopen Terminal."),
        ("MLflow page is blank", "Confirm mlflow ui is running from the project root with --backend-store-uri ./mlruns."),
    ]:
        cells = table.add_row().cells
        cells[0].text = problem
        cells[1].text = resolution

    section(d, "10 Submission Checklist")
    bullets(d, [
        "Run tests and linting with no failures.",
        "Capture EDA figures, MLflow comparison, API prediction, and Prometheus screenshots.",
        "Push the completed folder to your own GitHub repository.",
        "Replace the Kubernetes image placeholder and add a public URL or local access instructions.",
        "Use this guide while recording the 3 to 5 minute demonstration video.",
    ])
    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
